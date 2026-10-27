#!/usr/bin/env python3
"""
Rubric coding sheet, as Gordon (2026, 3.3): one row per topic of the whole-corpus model (NUM_TOPICS) and
the period models (PERIOD_TOPICS per period), for coding by hand on

    coherent   yes/no: the top words form a recognizable semantic field
    thematic   yes/no, coherent topics only: the field is recurring subject matter, not a register,
               procedure or paratext

Each row gives the topic's TOP_WORDS most probable words (by probability, not relevance: what the model
learned), the titles of its TOP_MOTIONS highest-loading motions (session, party of the first signatory, id)
and its party concentration, the counterpart of Gordon's author concentration: the share of the top
motions whose first signatory belongs to the most frequent party among them (1/TOP_MOTIONS to 1.0;
parties counted as in the topic explorer, so FP and L are one party; a motion without a party counts
for no party). Topic names and seed stability are left out, so they do not steer the judgement.

The coded sheet is read by analysis/rubric_summary.py.

Usage: python analysis/rubric_coding_sheet.py [--force]
Output: docs/rubric/RUBRIC_CODING_SHEET.xlsx (not overwritten without --force, since it holds the coding)
"""

import importlib
import pickle
import sys
from collections import Counter
from pathlib import Path

import numpy as np

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT / 'modeling'))

from lda_common import (  # noqa: E402
    BASELINE_MODEL_DIR, NUM_TOPICS, PERIOD_TOPICS, load_baseline_vectorizer, load_body_texts,
    load_period_body_texts, period_display, topic_word_distributions,
)

topic_explorer = importlib.import_module('14_generate_topic_explorer')

TOP_WORDS = 20
TOP_MOTIONS = 5
INPUT_DIR = PROJECT_ROOT / 'data' / 'output' / 'stopword_removed'
SHEET_PATH = PROJECT_ROOT / 'docs' / 'rubric' / 'RUBRIC_CODING_SHEET.xlsx'

COLUMNS = ['Model', 'Period', 'Topic', f'Top {TOP_WORDS} words', f'Top {TOP_MOTIONS} motions',
           'Party concentration', 'Most frequent party', 'Coherent', 'Thematic', 'Notes']
WHOLE_CORPUS = 'Whole corpus'
PERIOD_MODEL = 'Period'

INSTRUCTIONS = [
    ('Rubric', 'Gordon (2026, 3.3). Code every row once, before choosing chains for close reading.'),
    ('Coherent', 'yes / no: the top words form a recognizable semantic field.'),
    ('Order', 'Coherence first, as a rough filter; thematic content only for the coherent topics.'),
    ('Thematic', 'yes / no, only if coherent (leave empty otherwise): the field is recurring subject matter '
                 '(e.g. study support, preschool), not a register, procedure or paratext (budget language, '
                 'parliamentary formulas). Recurrence is judged within the model (for a period model, within '
                 'the period), not across periods.'),
    ('Motions', 'Session · party of the first signatory · id · title. Gordon judged from the top words and the '
                'top 5 texts themselves: open a motion by its id (e.g. in the topic explorer or at '
                'data.riksdagen.se) where the title is not enough.'),
    ('Party concentration', f'Computed: share of the top {TOP_MOTIONS} motions whose first signatory belongs to '
                            'the most frequent party among them (counterpart of Gordon\'s author concentration).'),
    ('Words', f'The {TOP_WORDS} most probable words (by probability, not the relevance ranking of the explorers).'),
    ('Not shown', 'Topic names and seed stability, so they do not steer the judgement.'),
    ('Notes', 'Optional: anything worth remembering for the report.'),
]


def party_concentration(parties):
    """Share of the motions whose party is the most frequent one among them; '?' (no party) counts for none.

    Returns (concentration, most frequent party or '' if no motion has a party).
    """
    known = Counter(p for p in parties if p != '?')
    if not known:
        return 1 / len(parties), ''
    party, count = max(known.items(), key=lambda item: (item[1], item[0]))
    return count / len(parties), party


def topic_rows(model_label, period_label, lda, vocab, theta, metadata, parties):
    """One sheet row per topic of one model."""
    rows = []
    topic_word = topic_word_distributions(lda)
    for k in range(lda.n_components):
        words = ', '.join(vocab[i] for i in np.argsort(topic_word[k])[::-1][:TOP_WORDS])
        top = np.argsort(theta[:, k])[::-1][:TOP_MOTIONS]
        codes = [parties.get(metadata[d]['doc_id'], '') for d in top]
        counted = [topic_explorer.party_as_counted(c) for c in codes]
        concentration, party = party_concentration(counted)
        motions = '\n'.join(
            f"{metadata[d].get('session', '')} · {topic_explorer.party_name(c)} · {metadata[d]['doc_id']} · "
            f"{' '.join(metadata[d].get('title', '').split())}"
            for d, c in zip(top, codes))
        rows.append([model_label, period_label, k + 1, words, motions, round(concentration, 2), party, '', '', ''])
    return rows


def normalized_theta(lda, doc_term):
    theta = lda.transform(doc_term)
    return theta / theta.sum(axis=1, keepdims=True)


def build_rows():
    parties = topic_explorer.load_first_signatory_parties(INPUT_DIR)
    model_dir = PROJECT_ROOT / 'models' / BASELINE_MODEL_DIR
    vectorizer = load_baseline_vectorizer(model_dir / 'baseline_vectorizer.pkl')
    with open(model_dir / f'baseline_lda_{NUM_TOPICS}_topics.pkl', 'rb') as f:
        lda = pickle.load(f)
    texts, metadata, _ = load_body_texts(INPUT_DIR)
    rows = topic_rows(WHOLE_CORPUS, '', lda, vectorizer.get_feature_names_out(),
                      normalized_theta(lda, vectorizer.transform(texts)), metadata, parties)

    periods, _ = load_period_body_texts(INPUT_DIR)
    period_dir = PROJECT_ROOT / 'models' / f'period_lda_{PERIOD_TOPICS}'
    for key, (period_texts, period_meta) in periods.items():
        with open(period_dir / f'period_{key}_lda.pkl', 'rb') as f:
            lda = pickle.load(f)
        with open(period_dir / f'period_{key}_vectorizer.pkl', 'rb') as f:
            vectorizer = pickle.load(f)
        label = period_display(key)
        rows += topic_rows(PERIOD_MODEL, label, lda, vectorizer.get_feature_names_out(),
                           normalized_theta(lda, vectorizer.transform(period_texts)), period_meta, parties)
    return rows


def write_sheet(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = 'Coding'
    ws.append(COLUMNS)
    for row in rows:
        ws.append(row)
    widths = {'A': 13, 'B': 16, 'C': 7, 'D': 48, 'E': 70, 'F': 13, 'G': 11, 'H': 10, 'I': 10, 'J': 30}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'D2'
    last = len(rows) + 1
    yes_no = DataValidation(type='list', formula1='"yes,no"', allow_blank=True)
    ws.add_data_validation(yes_no)
    yes_no.add(f'H2:I{last}')

    info = wb.create_sheet('Instructions')
    for item in INSTRUCTIONS:
        info.append(item)
    info.column_dimensions['A'].width = 20
    info.column_dimensions['B'].width = 100
    for row in info.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical='top')

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main():
    if SHEET_PATH.exists() and '--force' not in sys.argv[1:]:
        sys.exit(f'{SHEET_PATH} exists and may hold coding; pass --force to overwrite it.')
    rows = build_rows()
    write_sheet(rows, SHEET_PATH)
    print(f'✅ Coding sheet with {len(rows)} topics saved to: {SHEET_PATH}')


if __name__ == '__main__':
    main()
